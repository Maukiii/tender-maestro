from pathlib import Path
from typing import Callable

import docx
import fitz  # pymupdf
import openpyxl

from .config import CHUNK_SIZE


def chunk_text(text: str, max_words: int = CHUNK_SIZE) -> list[str]:
    words = text.split()
    return [
        " ".join(words[i : i + max_words])
        for i in range(0, len(words), max_words)
    ]


def parse_pdf(path: Path) -> list[str]:
    return chunk_text(extract_pdf_text(path))


def parse_docx(path: Path) -> list[str]:
    return chunk_text(extract_docx_text(path))


def parse_xlsx(path: Path) -> list[str]:
    return chunk_text(extract_xlsx_text(path))


def extract_pdf_text(path: Path) -> str:
    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n".join(parts)


def extract_docx_text(path: Path) -> str:
    doc = docx.Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def extract_xlsx_text(path: Path) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    blocks: list[str] = []
    for sheet in wb.worksheets:
        text = f"[{sheet.title}]\n"
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            text += row_text + "\n"
        blocks.append(text)
    return "\n".join(blocks)


def parse_text_file(path: Path) -> list[str]:
    return chunk_text(extract_text_file(path))


def extract_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8")


PARSERS: dict[str, Callable[[Path], list[str]]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".md": parse_text_file,
    ".txt": parse_text_file,
}

TEXT_EXTRACTORS: dict[str, Callable[[Path], str]] = {
    ".pdf": extract_pdf_text,
    ".docx": extract_docx_text,
    ".xlsx": extract_xlsx_text,
    ".md": extract_text_file,
    ".txt": extract_text_file,
}

